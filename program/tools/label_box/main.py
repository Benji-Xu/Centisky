"""
Excel 文件处理工具 - 集成标签箱唛功能
现代化设计 - 简洁·优雅·实用
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import os
import threading
from pathlib import Path
from datetime import datetime
import sys

# 添加父目录到路径以导入theme模块
sys.path.insert(0, str(Path(__file__).parent.parent.parent))
from theme import get_colors
from razer_ui import Razer3DCard, Razer3DRadio, Razer3DCheckbox
from unified_button import UnifiedButton
from theme_toggle import ThemeToggleButton

# 导入标签箱唛包装器
try:
    # 添加当前目录到Python路径
    current_dir = Path(__file__).parent
    if str(current_dir) not in sys.path:
        sys.path.insert(0, str(current_dir))
    
    import importlib
    import wrapper as wrapper_module
    importlib.reload(wrapper_module)
    from wrapper import process_excel_file
    LABEL_BOX_AVAILABLE = True
except Exception as e:
    LABEL_BOX_AVAILABLE = False
    print(f"警告：无法加载标签箱唛模块：{e}")
    import traceback
    traceback.print_exc()


class IntegratedApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Workit - 标签箱唛处理工具")
        self.root.geometry("1100x750")
        self.root.resizable(True, True)
        
        # 不设置窗口图标（用户不需要）
        
        # Razer风格配色 - 自动跟随系统深色/浅色模式
        self.colors = get_colors()
        
        self.root.configure(bg=self.colors['bg_main'])
        self.center_window()
        
        # 当前文件路径
        self.current_file = None
        self.processing = False
        
        # 输出路径 - 默认为系统下载文件夹
        self.output_dir = os.path.join(os.path.expanduser("~"), "Downloads")
        
        # 日志收集
        self.log_buffer = []
        
        # 处理选项（默认值）
        self.type_mode = tk.StringVar(value="auto")  # auto/3c/toy
        self.output_mode = tk.StringVar(value="both")  # both/label/box
        self.zip_mode = tk.BooleanVar(value=True)  # 是否打包zip（默认勾选）
        self.log_mode = tk.BooleanVar(value=True)  # 是否生成日志文件（默认勾选）
        
        # 店铺筛选（只影响标签生成）- 根据类型动态变化
        # 3C店铺
        self.shops_3c = ["外星人玩具", "三只梨", "兽", "兽无人机拆1", "兽无人机拆2"]
        # 玩具店铺
        self.shops_toy = ["外星人", "兽模型"]
        
        self.shop_filters = {}
        # 创建所有可能的店铺变量
        all_shops = set(self.shops_3c + self.shops_toy)
        for shop in all_shops:
            self.shop_filters[shop] = tk.BooleanVar(value=True)
        
        # 自动识别的类型（用于对比）
        self.detected_type = None
        
        # 创建样式
        self.setup_styles()
        
        # 创建界面
        self.create_widgets()
        
    def center_window(self):
        """窗口居中"""
        self.root.update_idletasks()
        width = 1100
        height = 750
        x = (self.root.winfo_screenwidth() // 2) - (width // 2)
        y = (self.root.winfo_screenheight() // 2) - (height // 2)
        self.root.geometry(f'{width}x{height}+{x}+{y}')
        
    def setup_styles(self):
        """设置样式"""
        style = ttk.Style()
        style.theme_use('clam')
        
        # 进度条样式
        style.configure(
            "Modern.Horizontal.TProgressbar",
            troughcolor=self.colors['bg_main'],
            background=self.colors['primary'],
            borderwidth=0,
            thickness=16
        )
        
    def create_widgets(self):
        """创建UI组件"""
        
        # 顶部标题区域（统一布局）
        header_frame = tk.Frame(self.root, bg=self.colors['bg_main'], height=80)
        header_frame.pack(fill=tk.X)
        header_frame.pack_propagate(False)

        # 左侧返回按钮
        back_btn = tk.Label(
            header_frame,
            text="< 返回首页",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_main'],
            fg=self.colors['text_muted'],
            cursor="hand2"
        )
        back_btn.place(relx=0.0, rely=0.61, x=40, anchor='w')
        back_btn.bind("<Button-1>", lambda e: self.back_to_launcher())
        back_btn.bind("<Enter>", lambda e: back_btn.config(fg=self.colors['text_primary']))
        back_btn.bind("<Leave>", lambda e: back_btn.config(fg=self.colors['text_muted']))

        # 右侧主题切换按钮（自定义日月图标）
        theme_btn = ThemeToggleButton(header_frame, command=self.toggle_theme)
        theme_btn.place(relx=1.0, rely=0.58, x=-40, anchor='e')

        # 右上角帮助按钮（?）
        help_btn = tk.Label(
            header_frame,
            text="?",
            font=("Microsoft YaHei UI", 13, "bold"),
            bg=self.colors['bg_main'],
            fg=self.colors['text_muted'],
            cursor="hand2"
        )
        help_btn.place(relx=1.0, rely=0.61, x=-80, anchor='e')
        help_btn.bind("<Button-1>", lambda e: self.open_help())
        help_btn.bind("<Enter>", lambda e: help_btn.config(fg=self.colors['text_primary']))
        help_btn.bind("<Leave>", lambda e: help_btn.config(fg=self.colors['text_muted']))

        # 中间标题
        title_container = tk.Frame(header_frame, bg=self.colors['bg_main'])
        title_container.place(relx=0.5, rely=0.61, anchor='center')
        
        title_label = tk.Label(
            title_container,
            text="标签箱唛处理工具",
            font=("Microsoft YaHei UI", 24, "bold"),
            bg=self.colors['bg_main'],
            fg=self.colors['text_primary']
        )
        title_label.pack()
        
        # 主内容区域
        content_frame = tk.Frame(self.root, bg=self.colors['bg_main'])
        content_frame.pack(fill=tk.BOTH, expand=True, padx=60, pady=35)
        
        # 文件选择卡片
        file_card_container, file_card = self.create_card(content_frame)
        file_card_container.pack(fill=tk.X, pady=(0, 18))
        
        # 标题行（包含输出路径按钮）
        title_row = tk.Frame(file_card, bg=self.colors['bg_card'])
        title_row.pack(fill=tk.X, pady=(0, 15))
        
        tk.Label(
            title_row,
            text="选择Excel文件",
            font=("Microsoft YaHei UI", 14, "bold"),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(side=tk.LEFT)
        
        # 右侧输出路径按钮（无背景，先pack齿轮再pack文件夹）
        output_btn_small = tk.Label(
            title_row,
            text="⚙",
            font=("Microsoft YaHei UI", 12),
            bg=self.colors['bg_card'],
            fg=self.colors['text_muted'],
            cursor="hand2"
        )
        output_btn_small.pack(side=tk.RIGHT)
        output_btn_small.bind("<Button-1>", lambda e: self.select_output_dir())
        output_btn_small.bind("<Enter>", lambda e: output_btn_small.config(fg=self.colors['primary']))
        output_btn_small.bind("<Leave>", lambda e: output_btn_small.config(fg=self.colors['text_muted']))
        
        self.output_path_label = tk.Label(
            title_row,
            text=f"📁 {os.path.basename(self.output_dir)}",
            font=("Microsoft YaHei UI", 9),
            bg=self.colors['bg_card'],
            fg=self.colors['text_muted']
        )
        self.output_path_label.pack(side=tk.RIGHT, padx=(0, 5))
        
        # 文件状态显示
        status_frame = tk.Frame(file_card, bg=self.colors['bg_main'])
        status_frame.pack(fill=tk.X, pady=(0, 15))
        
        self.file_label = tk.Label(
            status_frame,
            text="尚未选择文件",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_main'],
            fg=self.colors['text_muted'],
            anchor=tk.W,
            padx=15,
            pady=12
        )
        self.file_label.pack(fill=tk.X)
        
        # 按钮区域
        button_frame = tk.Frame(file_card, bg=self.colors['bg_card'])
        button_frame.pack(fill=tk.X)
        
        # 选择文件按钮（Razer 3D拟物化）
        self.select_btn = UnifiedButton(
            button_frame,
            text="选择 Excel 文件",
            command=self.select_file,
            style="primary",
            width=140,
            height=40
        )
        self.select_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        # 开始处理按钮（Razer 3D拟物化）
        self.process_btn = UnifiedButton(
            button_frame,
            text="开始处理",
            command=self.process_file,
            style="primary",
            width=120,
            height=40,
            state="disabled"
        )
        self.process_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        # 清除按钮（Razer 3D拟物化）
        self.clear_btn = UnifiedButton(
            button_frame,
            text="清除",
            command=self.clear_file,
            style="secondary",
            width=90,
            height=40,
            state="disabled"
        )
        self.clear_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        # 导入玩具标签按钮（居右显示）
        self.import_toy_btn = UnifiedButton(
            button_frame,
            text="导入玩具标签",
            command=lambda: self.import_pld_file_direct("玩具"),
            style="secondary",
            width=110,
            height=40
        )
        self.import_toy_btn.pack(side=tk.RIGHT, padx=(10, 0))
        
        # 导入3C标签按钮（居右显示）
        self.import_3c_btn = UnifiedButton(
            button_frame,
            text="导入3C标签",
            command=lambda: self.import_pld_file_direct("3C"),
            style="secondary",
            width=110,
            height=40
        )
        self.import_3c_btn.pack(side=tk.RIGHT)
        
        # 处理选项卡片
        options_card_container, options_card = self.create_card(content_frame)
        options_card_container.pack(fill=tk.X, pady=(0, 18))
        
        tk.Label(
            options_card,
            text="处理选项",
            font=("Microsoft YaHei UI", 14, "bold"),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(anchor=tk.W, pady=(0, 15))
        
        # 选项容器
        options_container = tk.Frame(options_card, bg=self.colors['bg_card'])
        options_container.pack(fill=tk.X)
        
        # 第一行：类型选择 + 店铺筛选
        row1 = tk.Frame(options_container, bg=self.colors['bg_card'])
        row1.pack(fill=tk.X, pady=(0, 15))
        
        # 类型选择
        type_frame = tk.Frame(row1, bg=self.colors['bg_card'])
        type_frame.pack(side=tk.LEFT, padx=(0, 30))
        
        tk.Label(
            type_frame,
            text="类型识别：",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(side=tk.LEFT, padx=(0, 10))
        
        self.create_flat_radio(type_frame, "自动", self.type_mode, "auto").pack(side=tk.LEFT, padx=(0, 8))
        self.create_flat_radio(type_frame, "3C", self.type_mode, "3c").pack(side=tk.LEFT, padx=(0, 8))
        self.create_flat_radio(type_frame, "玩具", self.type_mode, "toy").pack(side=tk.LEFT)
        
        # 店铺筛选（根据类型和输出模式动态显示）
        self.shop_filter_frame = tk.Frame(row1, bg=self.colors['bg_card'])
        self.shop_filter_frame.pack(side=tk.LEFT)
        
        tk.Label(
            self.shop_filter_frame,
            text="标签店铺：",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(side=tk.LEFT, padx=(0, 10))
        
        # 创建店铺复选框容器
        self.shop_checkboxes_frame = tk.Frame(self.shop_filter_frame, bg=self.colors['bg_card'])
        self.shop_checkboxes_frame.pack(side=tk.LEFT)
        
        # 第二行：输出内容 + 其他选项
        row2 = tk.Frame(options_container, bg=self.colors['bg_card'])
        row2.pack(fill=tk.X)
        
        # 输出内容选择
        output_frame = tk.Frame(row2, bg=self.colors['bg_card'])
        output_frame.pack(side=tk.LEFT, padx=(0, 30))
        
        tk.Label(
            output_frame,
            text="输出内容：",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(side=tk.LEFT, padx=(0, 10))
        
        self.create_flat_radio(output_frame, "标签+箱唛", self.output_mode, "both").pack(side=tk.LEFT, padx=(0, 8))
        self.create_flat_radio(output_frame, "仅标签", self.output_mode, "label").pack(side=tk.LEFT, padx=(0, 8))
        self.create_flat_radio(output_frame, "仅箱唛", self.output_mode, "box").pack(side=tk.LEFT, padx=(0, 8))
        self.create_flat_radio(output_frame, "仅预定表", self.output_mode, "reservation").pack(side=tk.LEFT)
        
        # 其他选项容器
        self.options_frame = tk.Frame(row2, bg=self.colors['bg_card'])
        self.options_frame.pack(side=tk.LEFT)
        
        # 创建复选框（保存引用以便动态显示/隐藏）
        self.zip_checkbox = self.create_flat_checkbox(self.options_frame, "打包为ZIP", self.zip_mode)
        self.zip_checkbox.pack(side=tk.LEFT, padx=(0, 8))
        
        self.log_checkbox = self.create_flat_checkbox(self.options_frame, "生成日志文件", self.log_mode)
        self.log_checkbox.pack(side=tk.LEFT)
        
        # 监听类型和输出模式变化，动态更新店铺筛选和其他选项显示
        self.type_mode.trace_add("write", lambda *args: self.update_shop_filters())
        self.output_mode.trace_add("write", lambda *args: self.update_ui_options())
        
        # 初始化店铺筛选和选项显示
        self.update_shop_filters()
        self.update_ui_options()
        
        # 螺旋桨管理卡片
        propeller_card_container, propeller_card = self.create_card(content_frame)
        propeller_card_container.pack(fill=tk.X, pady=(0, 18))
        
        # 螺旋桨管理标题行
        propeller_title_row = tk.Frame(propeller_card, bg=self.colors['bg_card'])
        propeller_title_row.pack(fill=tk.X, pady=(0, 15))
        
        tk.Label(
            propeller_title_row,
            text="螺旋桨映射管理",
            font=("Microsoft YaHei UI", 14, "bold"),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(side=tk.LEFT)
        
        
        # 螺旋桨管理内容
        propeller_content = tk.Frame(propeller_card, bg=self.colors['bg_card'])
        propeller_content.pack(fill=tk.X)
        
        # 第一行：商品编码输入
        prop_row1 = tk.Frame(propeller_content, bg=self.colors['bg_card'])
        prop_row1.pack(fill=tk.X, pady=(0, 10))
        
        tk.Label(
            prop_row1,
            text="商品编码：",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(side=tk.LEFT, padx=(0, 10))
        
        self.product_code_entry = tk.Entry(
            prop_row1,
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_main'],
            fg=self.colors['text_primary'],
            insertbackground=self.colors['text_primary'],
            relief=tk.FLAT,
            bd=1,
            width=20
        )
        self.product_code_entry.pack(side=tk.LEFT, padx=(0, 15))
        
        # PLD文件选择
        tk.Label(
            prop_row1,
            text="PLD文件：",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(side=tk.LEFT, padx=(0, 10))
        
        self.pld_file_label = tk.Label(
            prop_row1,
            text="未选择文件",
            font=("Microsoft YaHei UI", 9),
            bg=self.colors['bg_card'],
            fg=self.colors['text_muted']
        )
        self.pld_file_label.pack(side=tk.LEFT, padx=(0, 10))
        
        # 第二行：操作按钮
        prop_row2 = tk.Frame(propeller_content, bg=self.colors['bg_card'])
        prop_row2.pack(fill=tk.X)
        
        # 自动扫描按钮（放在最左边）
        self.scan_btn = UnifiedButton(
            prop_row2,
            text="自动扫描PLD",
            command=self.auto_scan_and_map,
            style="primary",
            width=120,
            height=32
        )
        self.scan_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        # 选择PLD文件按钮
        self.select_pld_btn = UnifiedButton(
            prop_row2,
            text="选择PLD文件",
            command=self.select_pld_file,
            style="secondary",
            width=120,
            height=32
        )
        self.select_pld_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        # 添加映射按钮
        self.add_mapping_btn = UnifiedButton(
            prop_row2,
            text="添加映射",
            command=self.add_propeller_mapping,
            style="secondary",
            width=100,
            height=32
        )
        self.add_mapping_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        # 查看映射按钮
        self.view_btn = UnifiedButton(
            prop_row2,
            text="查看映射",
            command=self.view_mappings,
            style="secondary",
            width=100,
            height=32
        )
        self.view_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        
        
        
        # 存储选中的PLD文件路径
        self.selected_pld_file = None
        
        # 启动时加载螺旋桨映射
        self.load_propeller_mappings()

        # 进度条卡片
        progress_card_container, progress_card = self.create_card(content_frame)
        progress_card_container.pack(fill=tk.X, pady=(0, 18))
        
        tk.Label(
            progress_card,
            text="处理进度",
            font=("Microsoft YaHei UI", 14, "bold"),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary']
        ).pack(anchor=tk.W, pady=(0, 12))
        
        self.progress_label = tk.Label(
            progress_card,
            text="等待开始...",
            font=("Microsoft YaHei UI", 10),
            bg=self.colors['bg_card'],
            fg=self.colors['text_muted'],
            anchor=tk.W
        )
        self.progress_label.pack(fill=tk.X, pady=(0, 8))
        
        self.progress_bar = ttk.Progressbar(
            progress_card,
            mode='determinate',
            maximum=100,
            style="Modern.Horizontal.TProgressbar",
            length=400
        )
        self.progress_bar.pack(fill=tk.X)
        self.progress_bar['value'] = 0
    
    def load_propeller_mappings(self):
        """启动时加载螺旋桨映射"""
        try:
            from propeller_config import load_mappings_from_file
            config_file = Path(__file__).parent / "propeller_mappings.json"
            
            if config_file.exists():
                success = load_mappings_from_file(str(config_file))
                if success:
                    print(f"成功加载螺旋桨映射文件：{config_file}")
                else:
                    print(f"加载螺旋桨映射文件失败：{config_file}")
            else:
                print(f"螺旋桨映射文件不存在：{config_file}")
                
        except Exception as e:
            print(f"加载螺旋桨映射时出错：{e}")
    
    
    def _extract_from_file_content(self, file_path):
        """从文件内容中提取商品编码"""
        try:
            print(f"从文件提取商品编码：{file_path}")
            
            # 尝试读取文件内容
            content = None
            for encoding in ['utf-8', 'gbk', 'latin1', 'cp1252']:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    print(f"成功读取文件，使用编码：{encoding}")
                    break
                except UnicodeDecodeError:
                    continue
                except Exception as e:
                    print(f"读取文件失败（{encoding}）：{e}")
                    continue
            
            if content is None:
                print("无法读取文件内容")
                return None
            
            print(f"文件内容长度：{len(content)} 字符")
            
            # 使用正则表达式提取商品编码
            import re
            
            patterns = [
                (r'商品编码[：:]\s*\*?(\d+)\*?', '明确标注的商品编码'),
                (r'SKU[：:]*\s*\*?(\d+)\*?', 'SKU编码'),
                (r'sku[：:]*\s*\*?(\d+)\*?', 'SKU编码(小写)'),
                (r'ID[：:]\s*\*?(\d+)\*?', 'ID编码'),
                (r'编码[：:]\s*\*?(\d+)\*?', '编码'),
                (r'\*(\d{12})\*', '星号包围的12位数字'),
                (r'\*(\d{11})\*', '星号包围的11位数字'),
                (r'\*(\d{10})\*', '星号包围的10位数字'),
                (r'(?:sku|SKU)\s+\*?(\d{10,})\*?', '条形码SKU'),
                (r'(\d{12})', '12位数字'),
                (r'(\d{11})', '11位数字'),
                (r'(\d{10})', '10位数字'),
                (r'(\d{9})', '9位数字'),
                (r'(\d{8})', '8位数字'),
            ]
            
            all_matches = []
            for pattern, description in patterns:
                matches = re.findall(pattern, content)
                for match in matches:
                    all_matches.append((match, description))
                    print(f"找到匹配（{description}）：{match}")
            
            if all_matches:
                best_match = max(all_matches, key=lambda x: len(x[0]))
                result = best_match[0]
                print(f"选择最佳匹配：{result} ({best_match[1]})")
                return result
            else:
                print("没有找到任何匹配")
                return None
                
        except Exception as e:
            print(f"从文件提取商品编码失败：{e}")
            return None
    
    def extract_product_code_from_pld(self, pld_filename, template_dir=None):
        """从PLD文件中提取商品编码"""
        try:
            print(f"开始提取商品编码：{pld_filename}")
            
            # 如果没有提供模板目录，尝试获取
            if template_dir is None:
                template_dir = self.get_template_directory()
            
            # 如果还是没有找到模板目录，返回None
            if not template_dir:
                print("未找到模板目录，无法提取商品编码")
                return None
            
            # 确保template_dir是Path对象
            if isinstance(template_dir, str):
                template_dir = Path(template_dir)
            
            print(f"在目录中搜索文件：{template_dir}")
            
            # 递归搜索匹配的PLD文件
            try:
                for pld_file in template_dir.rglob("*.pld"):
                    if pld_file.name == pld_filename:
                        print(f"精确匹配找到文件：{pld_file}")
                        return self._extract_from_file_content(pld_file)
                
                # 如果精确匹配失败，尝试模糊匹配
                for pld_file in template_dir.rglob("*.pld"):
                    # 检查文件名是否包含相同的关键词
                    if ("600" in pld_filename and "600" in pld_file.name and 
                        ("螺旋桨" in pld_file.name or "操旋奖" in pld_file.name)):
                        print(f"模糊匹配找到文件：{pld_file}")
                        return self._extract_from_file_content(pld_file)
                
                print(f"在目录中未找到匹配的文件：{pld_filename}")
                return None
                
            except Exception as e:
                print(f"搜索文件时出错：{e}")
                return None
            
        except Exception as e:
            print(f"提取商品编码失败：{e}")
            import traceback
            traceback.print_exc()
            return None
        
    def update_shop_filters(self):
        """根据类型和输出模式动态更新店铺筛选显示"""
        # 清空现有的复选框
        for widget in self.shop_checkboxes_frame.winfo_children():
            widget.destroy()
        
        # 判断是否需要显示店铺筛选
        type_val = self.type_mode.get()
        output_val = self.output_mode.get()
        
        # 只有在选择了具体类型（非自动）且输出包含标签时才显示（不包括仅预定表）
        if type_val in ["3c", "toy"] and output_val in ["both", "label"]:
            # 根据类型选择店铺列表
            shops = self.shops_3c if type_val == "3c" else self.shops_toy
            
            # 创建复选框
            for shop in shops:
                if shop in self.shop_filters:
                    self.create_flat_checkbox(self.shop_checkboxes_frame, shop, self.shop_filters[shop]).pack(side=tk.LEFT, padx=(0, 8))
            
            # 显示店铺筛选区域
            self.shop_filter_frame.pack(side=tk.LEFT)
        else:
            # 隐藏店铺筛选区域（包括选择仅预定表时）
            self.shop_filter_frame.pack_forget()
    
    def update_ui_options(self):
        """根据输出模式动态更新其他选项（ZIP、日志）的显示"""
        output_val = self.output_mode.get()
        
        # 当选择"仅预定表"时，隐藏ZIP和日志选项
        if output_val == "reservation":
            self.zip_checkbox.pack_forget()
            self.log_checkbox.pack_forget()
        else:
            # 其他模式显示这些选项
            self.zip_checkbox.pack(side=tk.LEFT, padx=(0, 8))
            self.log_checkbox.pack(side=tk.LEFT)
        
        # 同时更新店铺筛选（联动更新）
        self.update_shop_filters()
    
    def toggle_theme(self):
        """切换主题（保留数据）"""
        from theme import get_theme
        current_theme = get_theme()
        new_dark_mode = not current_theme.is_dark
        
        # 保存数据
        saved_file = self.current_file
        saved_output_dir = self.output_dir
        
        self.root.destroy()
        import theme
        theme._global_theme = theme.RazerTheme(dark_mode=new_dark_mode)
        
        new_root = tk.Tk()
        app = IntegratedApp(new_root)
        
        # 恢复数据
        if saved_file:
            app.current_file = saved_file
            app.output_dir = saved_output_dir
            app.file_label.config(
                text=f"✓ {Path(saved_file).name}",
                fg=app.colors['text_primary'],
                font=("Microsoft YaHei UI", 10, "bold")
            )
            if app.output_dir:
                app.output_path_label.config(text=f"📁 {os.path.basename(app.output_dir)}")
            app.process_btn.config_state("normal")
            app.clear_btn.config_state("normal")
        
        new_root.mainloop()
    
    def back_to_launcher(self):
        """返回首页"""
        try:
            # 导入launcher
            sys.path.insert(0, str(Path(__file__).parent.parent.parent))
            from launcher import ToolLauncher
            
            # 销毁当前窗口
            self.root.destroy()
            
            # 创建新窗口运行launcher
            new_root = tk.Tk()
            app = ToolLauncher(new_root)
            app.center_window()
            new_root.mainloop()
        except Exception as e:
            import traceback
            error_msg = f"返回首页失败：{e}\n\n{traceback.format_exc()}"
            print(error_msg)
            messagebox.showerror("返回失败", error_msg)
    
    def open_help(self):
        """显示使用说明（标签箱唛处理工具）"""
        try:
            from tkinter import Canvas, Frame

            doc_path = Path(__file__).parent / "标签箱唛处理工具使用说明.md"
            if not doc_path.exists():
                messagebox.showinfo("提示", f"未找到使用说明文件：\n{doc_path}")
                return

            raw = doc_path.read_text(encoding="utf-8", errors="ignore")

            help_win = tk.Toplevel(self.root)
            help_win.title("标签箱唛处理工具 - 使用说明")
            help_win.configure(bg=self.colors['bg_main'])

            help_win.update_idletasks()
            w, h = 800, 600
            sw = help_win.winfo_screenwidth()
            sh = help_win.winfo_screenheight()
            x = (sw // 2) - (w // 2)
            y = (sh // 2) - (h // 2)
            help_win.geometry(f"{w}x{h}+{x}+{y}")
            help_win.transient(self.root)
            help_win.grab_set()

            canvas = Canvas(help_win, bg=self.colors['bg_main'], highlightthickness=0)
            scrollbar = tk.Scrollbar(help_win, orient="vertical", command=canvas.yview)
            canvas.configure(yscrollcommand=scrollbar.set)
            scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
            canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

            inner = Frame(canvas, bg=self.colors['bg_main'])
            canvas.create_window((0, 0), window=inner, anchor='nw')

            def _on_configure(event):
                canvas.configure(scrollregion=canvas.bbox("all"))

            inner.bind("<Configure>", _on_configure)

            def _on_mousewheel(event):
                if event.delta:
                    canvas.yview_scroll(int(-event.delta / 120), "units")

            canvas.bind_all("<MouseWheel>", _on_mousewheel)

            import re
            lines = raw.splitlines()
            for line in lines:
                stripped = line.rstrip()

                # 分隔线 ---
                if stripped.strip() == "---":
                    tk.Label(inner, text="", bg=self.colors['bg_main']).pack(pady=2)
                    continue

                if not stripped.strip():
                    tk.Label(inner, text="", bg=self.colors['bg_main']).pack()
                    continue

                m = re.match(r"^\s*(#{1,6})\s+(.*)$", stripped)
                if m:
                    level = len(m.group(1))
                    text = m.group(2)
                    text = text.replace("**", "").replace("`", "")
                    size = 16 if level <= 2 else 12
                    weight = "bold"
                    tk.Label(
                        inner,
                        text=text,
                        font=("Microsoft YaHei UI", size, weight),
                        bg=self.colors['bg_main'],
                        fg=self.colors['text_primary'],
                        anchor='w',
                        justify='left',
                        wraplength=760
                    ).pack(fill=tk.X, padx=12, pady=(6 if level <= 2 else 4, 2))
                    continue

                m = re.match(r"^\s*[-*+]\s+(.*)$", stripped)
                if m:
                    text = m.group(1).replace("**", "").replace("`", "")
                    text = "• " + text
                    tk.Label(
                        inner,
                        text=text,
                        font=("Microsoft YaHei UI", 10),
                        bg=self.colors['bg_main'],
                        fg=self.colors['text_primary'],
                        anchor='w',
                        justify='left',
                        wraplength=760
                    ).pack(fill=tk.X, padx=24, pady=1)
                    continue

                text = stripped.replace("**", "").replace("`", "")
                tk.Label(
                    inner,
                    text=text,
                    font=("Microsoft YaHei UI", 10),
                    bg=self.colors['bg_main'],
                    fg=self.colors['text_primary'],
                    anchor='w',
                    justify='left',
                    wraplength=760
                ).pack(fill=tk.X, padx=12, pady=1)

        except Exception as e:
            messagebox.showerror("错误", f"无法打开使用说明：{e}")
    
    def create_card(self, parent):
        """创建Razer 3D拟物化卡片"""
        card_3d = Razer3DCard(parent)
        content = card_3d.get_content()
        
        content_padded = tk.Frame(content, bg=self.colors['bg_card'])
        content_padded.pack(fill=tk.BOTH, expand=True, padx=30, pady=22)
        
        return card_3d, content_padded
    
    def create_flat_radio(self, parent, text, variable, value):
        """创建Razer 3D拟物化单选框"""
        radio = Razer3DRadio(parent, text, variable, value)
        return radio
    
    def create_flat_checkbox(self, parent, text, variable):
        """创建Razer 3D拟物化复选框"""
        checkbox = Razer3DCheckbox(parent, text, variable)
        return checkbox
        
    def select_file(self):
        """选择Excel文件"""
        file_path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[
                ("Excel 文件", "*.xlsx *.xls *.xlsm"),
                ("所有文件", "*.*")
            ]
        )
        
        if file_path:
            self.current_file = file_path
            file_name = os.path.basename(file_path)
            self.file_label.config(
                text=f"✓ {file_name}",
                fg=self.colors['text_primary'],
                font=("Microsoft YaHei UI", 10, "bold")
            )
            
            # 自动更新输出目录为文件所在目录
            self.output_dir = os.path.dirname(file_path)
            self.output_path_label.config(text=f"📁 {os.path.basename(self.output_dir)}")
            
            # 启用处理和清除按钮
            self.process_btn.config_state("normal")
            self.clear_btn.config_state("normal")
            
            self.log_message(f"已选择文件：{file_name}")
            self.log_message(f"文件路径：{file_path}")
            self.log_message(f"输出路径已自动设置为：{self.output_dir}")
            
    def select_output_dir(self):
        """选择输出目录"""
        dir_path = filedialog.askdirectory(
            title="选择输出目录",
            initialdir=self.output_dir
        )
        
        if dir_path:
            self.output_dir = dir_path
            self.output_path_label.config(text=f"📁 {os.path.basename(self.output_dir)}")
            self.log_message(f"输出路径已更改为：{self.output_dir}")
    
    def clear_file(self):
        """清除当前文件"""
        self.current_file = None
        self.file_label.config(
            text="尚未选择文件",
            fg=self.colors['text_muted'],
            font=("Microsoft YaHei UI", 10)
        )
        
        self.process_btn.config_state("disabled")
        self.clear_btn.config_state("disabled")
        self.log_buffer.clear()  # 清空日志缓冲区
        self.progress_label.config(text="等待开始...")
        self.progress_bar['value'] = 0
    
    def log_message(self, message):
        """添加日志消息"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        log_line = f"[{timestamp}] {message}"
        
        # 收集日志到缓冲区
        self.log_buffer.append(log_line)
        
    def process_file(self):
        """处理Excel文件"""
        if not self.current_file or self.processing:
            return
            
        if not LABEL_BOX_AVAILABLE:
            messagebox.showerror(
                "错误",
                "标签箱唛模块未正确加载！\n请确保 标签箱唛.py 文件存在于同目录。"
            )
            return
        
        # 如果用户手动选择了类型，先快速检查是否与自动识别一致
        if self.type_mode.get() != "auto":
            # 快速读取检查类型
            try:
                from openpyxl import load_workbook
                wb = load_workbook(filename=self.current_file, data_only=True)
                
                # 导入核心模块快速识别
                from wrapper import core
                sheet_name_map = core.resolve_sheet_names(wb)
                b1_values = {exp: (core.read_b1(wb, real) if real else "") 
                             for exp, real in sheet_name_map.items()}
                auto_type = core.decide_label_type_by_b1(b1_values, sheet_name_map)
                
                user_type_map = {"3c": "3C", "toy": "玩具"}
                user_type = user_type_map.get(self.type_mode.get(), "")
                
                wb.close()
                
                # 如果类型不一致，弹窗让用户选择
                if auto_type != user_type:
                    result = messagebox.askyesnocancel(
                        "检测到类型不一致",
                        f"自动识别类型：{auto_type}\n"
                        f"您选择的类型：{user_type}\n\n"
                        f"是否继续使用您选择的类型({user_type})？\n\n"
                        f"【是】- 使用我选择的类型({user_type})\n"
                        f"【否】- 使用自动识别类型({auto_type})\n"
                        f"【取消】- 取消处理"
                    )
                    
                    if result is None:  # 取消
                        return
                    elif result is False:  # 使用自动识别
                        self.type_mode.set("auto")
                        self.log_message(f"已切换为自动识别类型：{auto_type}")
                    # result is True 则继续使用用户选择的类型
                    
            except Exception as e:
                self.log_message(f"类型检查失败：{e}")
            
        # 在新线程中执行处理
        self.processing = True
        self.select_btn.config(state=tk.DISABLED)
        self.process_btn.config_state("disabled")
        self.clear_btn.config_state("disabled")
        
        self.progress_label.config(text="正在处理中...")
        self.progress_bar['value'] = 0
        
        thread = threading.Thread(target=self._do_process, daemon=True)
        thread.start()
        
    def update_progress(self, value, text=None):
        """更新进度条"""
        self.progress_bar['value'] = value
        if text:
            self.progress_label.config(text=text)
        self.root.update()
    
    def _do_process(self):
        """实际处理逻辑（在后台线程执行）"""
        try:
            # 清空日志缓冲区
            self.log_buffer.clear()
            
            self.log_message("="*50)
            self.log_message("开始处理...")
            self.log_message(f"工作簿：{os.path.basename(self.current_file)}")
            self.log_message(f"输出路径：{self.output_dir}")
            self.log_message("")
            
            # 显示选项
            type_mode_text = {"auto": "自动识别", "3c": "强制3C", "toy": "强制玩具"}[self.type_mode.get()]
            output_mode_text = {"both": "标签+箱唛", "label": "仅标签", "box": "仅箱唛", "reservation": "仅预定表"}[self.output_mode.get()]
            self.log_message(f"类型模式：{type_mode_text}")
            self.log_message(f"输出模式：{output_mode_text}")
            if self.zip_mode.get():
                self.log_message("打包ZIP：是")
            if self.log_mode.get():
                self.log_message("生成日志：是")
            self.log_message("")
            
            self.update_progress(10, "正在读取Excel文件...")
            
            # 获取选中的店铺
            selected_shops = [shop for shop, var in self.shop_filters.items() if var.get()]
            
            # 重新加载wrapper模块，确保获取最新的代码
            import importlib
            import wrapper as wrapper_module
            importlib.reload(wrapper_module)
            from wrapper import process_excel_file as process_excel_file_reloaded
            
            # 调用包装器函数
            result = process_excel_file_reloaded(
                self.current_file, 
                output_base=self.output_dir, 
                callback=self.log_message,
                progress_callback=self.update_progress,
                type_mode=self.type_mode.get(),
                output_mode=self.output_mode.get(),
                create_zip=self.zip_mode.get(),
                save_log=self.log_mode.get(),
                selected_shops=selected_shops
            )
            
            # 检查类型不匹配 - 不再需要，改为事前确认
            pass
            
            if result.get("success"):
                self.root.after(0, lambda: self.progress_label.config(text="✓ 处理完成！", fg=self.colors['success']))
                self.root.after(0, lambda: self.progress_bar.__setitem__('value', 100))
                self.log_message("")
                self.log_message("="*50)
                self.log_message("✓ 处理成功完成！")
                
                # 如果勾选了保存日志，将日志写入文件（保存在输出目录根目录）
                if self.log_mode.get() and result.get('mmdd') and result.get('label_type_name'):
                    try:
                        # 日志文件名：1027-3C标签箱唛-处理日志.txt
                        log_filename = f"{result['mmdd']}-{result['label_type_name']}标签箱唛-处理日志.txt"
                        log_file_path = os.path.join(self.output_dir, log_filename)
                        with open(log_file_path, 'w', encoding='utf-8') as f:
                            f.write("\n".join(self.log_buffer))
                        self.log_message(f"\n日志已保存：{log_file_path}")
                    except Exception as log_err:
                        self.log_message(f"\n保存日志失败：{log_err}")
                
                # 构建完成消息
                msg_lines = []
                
                # 检查是否为仅预定表模式
                if result.get('reservation_only'):
                    msg_lines.append(f"标签类型：{result.get('label_type', 'N/A')}")
                    msg_lines.append(f"日期标识：{result.get('mmdd', 'N/A')}")
                    msg_lines.append(f"预定表行数：{result.get('total_rows', 0)}")
                    msg_lines.append(f"\n输出文件：{result.get('output_path', 'N/A')}")
                else:
                    msg_lines.append(f"标签类型：{result.get('label_type', 'N/A')}")
                    msg_lines.append(f"日期标识：{result.get('mmdd', 'N/A')}")
                    
                    # 添加标签统计信息
                    total_expected = result.get('total_expected', 0)
                    total_copied = result.get('total_copied', 0)
                    total_missing = result.get('total_missing', 0)
                    
                    msg_lines.append(f"\n标签统计：")
                    msg_lines.append(f"  应该生成：{total_expected} 个")
                    msg_lines.append(f"  已生成：{total_copied} 个")
                    if total_missing > 0:
                        msg_lines.append(f"  缺少：{total_missing} 个")
                    
                    if result.get('box_ok') or result.get('box_warn'):
                        msg_lines.append(f"\n箱唛统计：")
                        msg_lines.append(f"  成功：{result.get('box_ok', 0)} 个")
                        if result.get('box_warn', 0) > 0:
                            msg_lines.append(f"  警告：{result.get('box_warn', 0)} 个")
                    
                    # 输出目录
                    msg_lines.append("\n输出文件夹：")
                    if result.get('label_output'):
                        msg_lines.append(f"  标签：{result['label_output']}")
                    if result.get('box_output'):
                        msg_lines.append(f"  箱唛：{result['box_output']}")
                    if result.get('main_output'):
                        msg_lines.append(f"\n输出位置：{result['main_output']}")
                    
                    # 如果有缺少标签报告，显示报告文件路径
                    if result.get('missing_report_path'):
                        msg_lines.append(f"\n缺少标签报告：{result['missing_report_path']}")
                
                for line in msg_lines:
                    self.log_message(line)
                
                # 检查是否有缺少的标签
                total_missing = result.get("total_missing", 0)
                self.log_message(f"\n[DEBUG] total_missing: {total_missing}")
                if total_missing > 0:
                    missing_details = result.get("missing_details", {})
                    self.log_message(f"[DEBUG] missing_details: {missing_details}")
                    missing_lines = [
                        f"缺少 {total_missing} 个标签，分别是：\n"
                    ]
                    for sheet, details_list in missing_details.items():
                        if details_list:
                            missing_lines.append(f"\n{sheet}：")
                            for detail in details_list:
                                sku = detail.get('sku', '')
                                e_val = detail.get('e_val', '')
                                # 显示商品编号和E列值
                                if sku:
                                    if e_val and e_val != 'None':
                                        missing_lines.append(f"  • {sku} ({e_val})")
                                    else:
                                        missing_lines.append(f"  • {sku}")
                                else:
                                    missing_lines.append(f"  • (未知)")
                    
                    missing_msg = "".join(missing_lines)
                    self.log_message("\n" + missing_msg)
                    
                    # 显示缺少标签的警告对话框
                    self.root.after(0, lambda msg=missing_msg: messagebox.showwarning(
                        "警告：缺少标签",
                        msg + "\n\n请检查模板目录中是否存在对应的PLD文件"
                    ))
                
                # 自动导出统计报告
                try:
                    from pathlib import Path
                    auto_export_path = Path(__file__).parent / "auto_export.py"
                    import importlib.util
                    spec = importlib.util.spec_from_file_location("auto_export", auto_export_path)
                    auto_export = importlib.util.module_from_spec(spec)
                    spec.loader.exec_module(auto_export)
                    auto_export.auto_export_after_process(
                        result.get('main_output'),
                        result.get('mmdd'),
                        result.get('total_expected', 0),
                        result.get('total_copied', 0),
                        result.get('missing_map', {})
                    )
                except Exception as e:
                    self.log_message(f"\n[导出] 自动导出失败：{e}")
                
                # 显示完成对话框
                self.root.after(0, lambda: messagebox.showinfo(
                    "处理完成",
                    "\n".join(msg_lines)
                ))
                
            else:
                error_msg = result.get("error", "未知错误")
                self.root.after(0, lambda: self.progress_label.config(text="✗ 处理失败", fg=self.colors['danger']))
                self.log_message("")
                self.log_message(f"✗ 处理失败：{error_msg}")
                
                if result.get("traceback"):
                    self.log_message("\n详细错误信息：")
                    self.log_message(result["traceback"])
                
                self.root.after(0, lambda: messagebox.showerror("处理失败", error_msg))
                
        except Exception as e:
            error_msg = str(e)
            self.log_message(f"\n✗ 发生异常：{error_msg}")
            import traceback
            error_trace = traceback.format_exc()
            self.log_message(error_trace)
            
            self.root.after(0, lambda: self.progress_label.config(text="✗ 发生异常", fg=self.colors['danger']))
            self.root.after(0, lambda msg=error_msg: messagebox.showerror("异常", f"发生异常：{msg}"))
            
        finally:
            # 恢复界面状态
            self.processing = False
            self.root.after(0, lambda: self.select_btn.config_state("normal"))
            self.root.after(0, lambda: self.process_btn.config_state("normal"))
            self.root.after(0, lambda: self.clear_btn.config_state("normal"))
    
    def select_pld_file(self):
        """选择PLD文件"""
        try:
            file_path = filedialog.askopenfilename(
                title="选择螺旋桨PLD文件",
                filetypes=[
                    ("PLD文件", "*.pld"),
                    ("所有文件", "*.*")
                ],
                initialdir=os.path.expanduser("~")
            )
            
            if file_path:
                self.selected_pld_file = file_path
                filename = os.path.basename(file_path)
                self.pld_file_label.config(
                    text=filename,
                    fg=self.colors['text_primary']
                )
                self.log_message(f"已选择PLD文件：{filename}")
            
        except Exception as e:
            messagebox.showerror("错误", f"选择文件失败：{e}")
    
    def add_propeller_mapping(self):
        """添加螺旋桨映射"""
        try:
            # 获取商品编码
            product_code = self.product_code_entry.get().strip()
            if not product_code:
                messagebox.showwarning("提示", "请输入商品编码")
                return
            
            # 检查是否选择了PLD文件
            if not self.selected_pld_file:
                messagebox.showwarning("提示", "请先选择PLD文件")
                return
            
            # 导入螺旋桨配置模块
            try:
                from propeller_config import add_propeller_mapping, save_mappings_to_file, get_all_mappings
            except ImportError:
                messagebox.showerror("错误", "无法导入螺旋桨配置模块")
                return
            
            # 获取PLD文件名
            pld_filename = os.path.basename(self.selected_pld_file)
            
            # 检查是否已存在映射
            existing_mappings = get_all_mappings()
            if product_code in existing_mappings:
                result = messagebox.askyesno(
                    "确认覆盖", 
                    f"商品编码 {product_code} 已映射到 {existing_mappings[product_code]}\n\n是否覆盖为 {pld_filename}？"
                )
                if not result:
                    return
            
            # 复制PLD文件到模板目录（如果需要）
            template_dir = self.get_template_directory()
            if template_dir:
                target_path = template_dir / pld_filename
                if not target_path.exists():
                    import shutil
                    shutil.copy2(self.selected_pld_file, target_path)
                    self.log_message(f"已复制PLD文件到模板目录：{target_path}")
            
            # 添加映射
            add_propeller_mapping(product_code, pld_filename)
            
            # 保存到配置文件
            config_file = Path(__file__).parent / "propeller_mappings.json"
            save_mappings_to_file(str(config_file))
            
            # 清空输入
            self.product_code_entry.delete(0, tk.END)
            self.selected_pld_file = None
            self.pld_file_label.config(
                text="未选择文件",
                fg=self.colors['text_muted']
            )
            
            self.log_message(f"✓ 已添加螺旋桨映射：{product_code} -> {pld_filename}")
            messagebox.showinfo("成功", f"已添加映射：\n{product_code} -> {pld_filename}")
            
        except Exception as e:
            self.log_message(f"✗ 添加映射失败：{e}")
            messagebox.showerror("错误", f"添加映射失败：{e}")
    
    def get_template_directory(self):
        """获取模板目录路径"""
        try:
            # 尝试多个可能的模板目录位置
            possible_dirs = [
                Path(__file__).parent.parent.parent.parent / "templates" / "标签模板",
                Path(__file__).parent.parent.parent / "templates" / "标签模板", 
                Path(__file__).parent / "templates" / "标签模板",
                Path(__file__).parent / "标签模板"
            ]
            
            for template_dir in possible_dirs:
                if template_dir.exists():
                    return template_dir
            
            return None
        except Exception:
            return None
    
    def auto_scan_and_map(self):
        """自动扫描PLD文件并提供映射界面"""
        try:
            template_dir = self.get_template_directory()
            if not template_dir:
                # 让用户选择模板目录
                template_dir = filedialog.askdirectory(title="选择模板目录")
                if not template_dir:
                    return
                template_dir = Path(template_dir)
            
            # 扫描螺旋桨相关的PLD文件
            propeller_files = []
            for pld_file in template_dir.rglob("*.pld"):
                filename = pld_file.name
                if any(keyword in filename for keyword in ["螺旋桨", "propeller", "螺桨"]):
                    propeller_files.append(filename)
            
            if not propeller_files:
                messagebox.showinfo("扫描结果", "未发现螺旋桨相关的PLD文件")
                return
            
            # 检查未映射的文件
            try:
                from propeller_config import get_all_mappings
                mapped_files = set(get_all_mappings().values())
                unmapped_files = [f for f in propeller_files if f not in mapped_files]
            except ImportError:
                unmapped_files = propeller_files
            
            if not unmapped_files:
                messagebox.showinfo("扫描结果", f"扫描完成！\n\n发现 {len(propeller_files)} 个螺旋桨文件，全部已映射")
                return
            
            # 打开映射窗口，传递模板目录
            self.open_mapping_window(unmapped_files, template_dir)
            self.log_message(f"自动扫描完成：发现 {len(propeller_files)} 个螺旋桨文件，{len(unmapped_files)} 个未映射")
            
        except Exception as e:
            self.log_message(f"✗ 自动扫描失败：{e}")
            messagebox.showerror("错误", f"自动扫描失败：{e}")
    
    def open_mapping_window(self, unmapped_files, template_dir=None):
        """打开映射窗口，允许用户为未映射的文件添加商品编码"""
        # 创建映射窗口
        mapping_window = tk.Toplevel(self.root)
        mapping_window.title("添加螺旋桨映射")
        mapping_window.geometry("900x600")
        mapping_window.configure(bg=self.colors['bg_main'])
        
        # 居中显示
        mapping_window.transient(self.root)
        mapping_window.grab_set()
        
        # 窗口居中
        mapping_window.update_idletasks()
        x = (mapping_window.winfo_screenwidth() // 2) - 450
        y = (mapping_window.winfo_screenheight() // 2) - 300
        mapping_window.geometry(f'900x600+{x}+{y}')
        
        # 顶部区域
        header_frame = tk.Frame(mapping_window, bg=self.colors['bg_main'])
        header_frame.pack(fill=tk.X, padx=30, pady=(30, 20))
        
        # 标题
        title_label = tk.Label(
            header_frame,
            text=f"发现 {len(unmapped_files)} 个未映射的螺旋桨文件",
            font=("Microsoft YaHei UI", 18, "bold"),
            bg=self.colors['bg_main'],
            fg=self.colors['text_primary']
        )
        title_label.pack()
        
        # 说明文字
        info_label = tk.Label(
            header_frame,
            text="系统已自动提取商品编码作为提示，请确认或修改（留空的文件将不会添加映射）",
            font=("Microsoft YaHei UI", 11),
            bg=self.colors['bg_main'],
            fg=self.colors['text_muted']
        )
        info_label.pack(pady=(10, 0))
        
        # 主内容区域
        content_frame = tk.Frame(mapping_window, bg=self.colors['bg_main'])
        content_frame.pack(fill=tk.BOTH, expand=True, padx=30, pady=(0, 20))
        
        # 表头
        header_card = tk.Frame(content_frame, bg=self.colors['bg_card'], relief=tk.FLAT, bd=1)
        header_card.pack(fill=tk.X, pady=(0, 5))
        
        header_inner = tk.Frame(header_card, bg=self.colors['bg_card'])
        header_inner.pack(fill=tk.X, padx=20, pady=12)
        
        # 使用Grid布局来确保表头对齐
        header_inner.grid_columnconfigure(0, weight=1)
        header_inner.grid_columnconfigure(1, weight=0, minsize=200)
        
        tk.Label(
            header_inner,
            text="PLD文件名",
            font=("Microsoft YaHei UI", 12, "bold"),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary'],
            anchor=tk.W
        ).grid(row=0, column=0, sticky="ew", padx=(0, 20))
        
        tk.Label(
            header_inner,
            text="商品编码",
            font=("Microsoft YaHei UI", 12, "bold"),
            bg=self.colors['bg_card'],
            fg=self.colors['text_primary'],
            anchor=tk.CENTER
        ).grid(row=0, column=1, sticky="ew")
        
        # 创建一个固定高度的列表容器，避免滚动条问题
        list_container = tk.Frame(content_frame, bg=self.colors['bg_main'])
        list_container.pack(fill=tk.BOTH, expand=True)
        
        # 如果文件数量较少，使用固定布局；如果较多，使用滚动布局
        if len(unmapped_files) <= 8:
            # 直接布局，不需要滚动
            scrollable_frame = list_container
            entry_widgets = {}
            
            # 为每个文件创建输入行
            for i, filename in enumerate(unmapped_files):
                # 创建行卡片
                row_card = tk.Frame(
                    scrollable_frame,
                    bg=self.colors['bg_card'],
                    relief=tk.FLAT,
                    bd=1,
                    height=50
                )
                row_card.pack(fill=tk.X, pady=2)
                row_card.pack_propagate(False)  # 固定高度
                
                # 使用Grid布局确保对齐
                row_card.grid_columnconfigure(0, weight=1)
                row_card.grid_columnconfigure(1, weight=0, minsize=200)
                
                # 文件名标签（左侧）
                file_label = tk.Label(
                    row_card,
                    text=filename,
                    font=("Microsoft YaHei UI", 11),
                    bg=self.colors['bg_card'],
                    fg=self.colors['text_primary'],
                    anchor=tk.W
                )
                file_label.grid(row=0, column=0, sticky="ew", padx=20, pady=10)
                
                # 尝试提取商品编码（传递模板目录）
                extracted_code = self.extract_product_code_from_pld(filename, template_dir)
                
                # 商品编码输入框
                code_entry = tk.Entry(
                    row_card,
                    font=("Microsoft YaHei UI", 11),
                    bg=self.colors['bg_main'],
                    fg=self.colors['text_primary'],
                    insertbackground=self.colors['text_primary'],
                    relief=tk.FLAT,
                    bd=1,
                    width=20,
                    justify=tk.CENTER
                )
                code_entry.grid(row=0, column=1, sticky="ew", padx=20, pady=10)
                
                # 如果提取到商品编码，设置为默认值
                if extracted_code:
                    code_entry.insert(0, extracted_code)
                    # 设置为灰色文本表示这是提取的值
                    code_entry.config(fg=self.colors['text_muted'])
                
                # 添加边框效果和焦点处理
                def on_focus_in(event, entry=code_entry, original_code=extracted_code):
                    entry.config(bg=self.colors['bg_card'], relief=tk.SOLID, bd=1, fg=self.colors['text_primary'])
                    # 如果内容是提取的原始值，选中全部文本方便修改
                    if original_code and entry.get() == original_code:
                        entry.select_range(0, tk.END)
                
                def on_focus_out(event, entry=code_entry, original_code=extracted_code):
                    entry.config(bg=self.colors['bg_main'], relief=tk.FLAT, bd=1)
                    # 如果输入框为空且有提取的编码，恢复提取的编码
                    if not entry.get().strip() and original_code:
                        entry.insert(0, original_code)
                        entry.config(fg=self.colors['text_muted'])
                    elif entry.get().strip():
                        entry.config(fg=self.colors['text_primary'])
                
                code_entry.bind("<FocusIn>", on_focus_in)
                code_entry.bind("<FocusOut>", on_focus_out)
                
                # 存储输入框引用
                entry_widgets[filename] = code_entry
        
        else:
            # 使用滚动布局
            canvas = tk.Canvas(list_container, bg=self.colors['bg_main'], highlightthickness=0)
            scrollbar = ttk.Scrollbar(list_container, orient="vertical", command=canvas.yview)
            scrollable_frame = tk.Frame(canvas, bg=self.colors['bg_main'])
            
            scrollable_frame.bind(
                "<Configure>",
                lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
            )
            
            canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
            canvas.configure(yscrollcommand=scrollbar.set)
            
            # 存储输入框引用
            entry_widgets = {}
            
            # 为每个文件创建输入行
            for i, filename in enumerate(unmapped_files):
                # 创建行卡片
                row_card = tk.Frame(
                    scrollable_frame,
                    bg=self.colors['bg_card'],
                    relief=tk.FLAT,
                    bd=1,
                    height=50
                )
                row_card.pack(fill=tk.X, pady=2)
                row_card.pack_propagate(False)  # 固定高度
                
                # 使用Grid布局确保对齐
                row_card.grid_columnconfigure(0, weight=1)
                row_card.grid_columnconfigure(1, weight=0, minsize=200)
                
                # 文件名标签（左侧）
                file_label = tk.Label(
                    row_card,
                    text=filename,
                    font=("Microsoft YaHei UI", 11),
                    bg=self.colors['bg_card'],
                    fg=self.colors['text_primary'],
                    anchor=tk.W
                )
                file_label.grid(row=0, column=0, sticky="ew", padx=20, pady=10)
                
                # 尝试提取商品编码（传递模板目录）
                extracted_code = self.extract_product_code_from_pld(filename, template_dir)
                
                # 商品编码输入框
                code_entry = tk.Entry(
                    row_card,
                    font=("Microsoft YaHei UI", 11),
                    bg=self.colors['bg_main'],
                    fg=self.colors['text_primary'],
                    insertbackground=self.colors['text_primary'],
                    relief=tk.FLAT,
                    bd=1,
                    width=20,
                    justify=tk.CENTER
                )
                code_entry.grid(row=0, column=1, sticky="ew", padx=20, pady=10)
                
                # 如果提取到商品编码，设置为默认值
                if extracted_code:
                    code_entry.insert(0, extracted_code)
                    # 设置为灰色文本表示这是提取的值
                    code_entry.config(fg=self.colors['text_muted'])
                
                # 添加边框效果和焦点处理
                def on_focus_in(event, entry=code_entry, original_code=extracted_code):
                    entry.config(bg=self.colors['bg_card'], relief=tk.SOLID, bd=1, fg=self.colors['text_primary'])
                    # 如果内容是提取的原始值，选中全部文本方便修改
                    if original_code and entry.get() == original_code:
                        entry.select_range(0, tk.END)
                
                def on_focus_out(event, entry=code_entry, original_code=extracted_code):
                    entry.config(bg=self.colors['bg_main'], relief=tk.FLAT, bd=1)
                    # 如果输入框为空且有提取的编码，恢复提取的编码
                    if not entry.get().strip() and original_code:
                        entry.insert(0, original_code)
                        entry.config(fg=self.colors['text_muted'])
                    elif entry.get().strip():
                        entry.config(fg=self.colors['text_primary'])
                
                code_entry.bind("<FocusIn>", on_focus_in)
                code_entry.bind("<FocusOut>", on_focus_out)
                
                # 存储输入框引用
                entry_widgets[filename] = code_entry
            
            canvas.pack(side="left", fill="both", expand=True)
            scrollbar.pack(side="right", fill="y")
            
            # 绑定鼠标滚轮事件
            def _on_mousewheel(event):
                canvas.yview_scroll(int(-1*(event.delta/120)), "units")
            
            canvas.bind_all("<MouseWheel>", _on_mousewheel)
            
            # 存储canvas引用以便后续解绑
            mapping_window.canvas = canvas
        
        # 底部按钮区域
        button_frame = tk.Frame(mapping_window, bg=self.colors['bg_main'])
        button_frame.pack(fill=tk.X, padx=30, pady=(0, 30))
        
        # 按钮容器
        button_container = tk.Frame(button_frame, bg=self.colors['bg_main'])
        button_container.pack()
        
        def save_mappings():
            """保存所有映射"""
            added_count = 0
            try:
                from propeller_config import add_propeller_mapping, save_mappings_to_file
                
                for filename, entry in entry_widgets.items():
                    code = entry.get().strip()
                    if code:  # 只保存非空的映射
                        add_propeller_mapping(code, filename)
                        added_count += 1
                        self.log_message(f"✓ 已添加映射：{code} -> {filename}")
                
                if added_count > 0:
                    # 保存到配置文件
                    config_file = Path(__file__).parent / "propeller_mappings.json"
                    save_mappings_to_file(str(config_file))
                    
                    messagebox.showinfo("成功", f"已添加 {added_count} 个映射")
                    mapping_window.destroy()
                else:
                    messagebox.showwarning("提示", "没有输入任何商品编码")
                    
            except Exception as e:
                messagebox.showerror("错误", f"保存映射失败：{e}")
        
        def skip_all():
            """跳过所有映射"""
            mapping_window.destroy()
        
        def clear_all():
            """清空所有输入"""
            for entry in entry_widgets.values():
                entry.delete(0, tk.END)
        
        # 保存按钮
        save_btn = UnifiedButton(
            button_container,
            text=f"保存映射 ({len(unmapped_files)})",
            command=save_mappings,
            style="primary",
            width=160,
            height=40
        )
        save_btn.pack(side=tk.LEFT, padx=(0, 15))
        
        # 清空按钮
        clear_btn = UnifiedButton(
            button_container,
            text="清空输入",
            command=clear_all,
            style="secondary",
            width=120,
            height=40
        )
        clear_btn.pack(side=tk.LEFT, padx=(0, 15))
        
        # 跳过按钮
        skip_btn = UnifiedButton(
            button_container,
            text="跳过",
            command=skip_all,
            style="secondary",
            width=100,
            height=40
        )
        skip_btn.pack(side=tk.LEFT)
        
        # 当窗口关闭时解绑事件
        def on_closing():
            # 如果有canvas，解绑滚轮事件
            if hasattr(mapping_window, 'canvas'):
                mapping_window.canvas.unbind_all("<MouseWheel>")
            mapping_window.destroy()
        
        mapping_window.protocol("WM_DELETE_WINDOW", on_closing)
    
    def show_propeller_mappings(self):
        """显示所有螺旋桨映射"""
        try:
            from propeller_config import get_all_mappings
            mappings = get_all_mappings()
            
            if not mappings:
                messagebox.showinfo("映射列表", "当前没有螺旋桨映射")
                return
            
            # 创建映射显示窗口
            mapping_window = tk.Toplevel(self.root)
            mapping_window.title("螺旋桨映射列表")
            mapping_window.geometry("800x500")
            mapping_window.configure(bg=self.colors['bg_main'])
            
            # 居中显示
            mapping_window.transient(self.root)
            mapping_window.grab_set()
            
            # 标题
            title_label = tk.Label(
                mapping_window,
                text="螺旋桨映射列表",
                font=("Microsoft YaHei UI", 16, "bold"),
                bg=self.colors['bg_main'],
                fg=self.colors['text_primary']
            )
            title_label.pack(pady=20)
            
            # 创建表格框架
            table_frame = tk.Frame(mapping_window, bg=self.colors['bg_main'])
            table_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 20))
            
            # 表格标题
            header_frame = tk.Frame(table_frame, bg=self.colors['bg_card'], height=40)
            header_frame.pack(fill=tk.X, pady=(0, 2))
            header_frame.pack_propagate(False)
            
            tk.Label(
                header_frame,
                text="商品编码",
                font=("Microsoft YaHei UI", 11, "bold"),
                bg=self.colors['bg_card'],
                fg=self.colors['text_primary']
            ).pack(side=tk.LEFT, padx=20, pady=10)
            
            tk.Label(
                header_frame,
                text="PLD文件名",
                font=("Microsoft YaHei UI", 11, "bold"),
                bg=self.colors['bg_card'],
                fg=self.colors['text_primary']
            ).pack(side=tk.LEFT, padx=100, pady=10)
            
            # 滚动区域
            canvas = tk.Canvas(table_frame, bg=self.colors['bg_main'])
            scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=canvas.yview)
            scrollable_frame = tk.Frame(canvas, bg=self.colors['bg_main'])
            
            scrollable_frame.bind(
                "<Configure>",
                lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
            )
            
            canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
            canvas.configure(yscrollcommand=scrollbar.set)
            
            # 添加映射行
            for i, (code, filename) in enumerate(mappings.items()):
                row_frame = tk.Frame(
                    scrollable_frame, 
                    bg=self.colors['bg_card'] if i % 2 == 0 else self.colors['bg_main'],
                    height=35
                )
                row_frame.pack(fill=tk.X, pady=1)
                row_frame.pack_propagate(False)
                
                tk.Label(
                    row_frame,
                    text=code,
                    font=("Microsoft YaHei UI", 10),
                    bg=row_frame['bg'],
                    fg=self.colors['text_primary'],
                    anchor=tk.W
                ).pack(side=tk.LEFT, padx=20, pady=5, fill=tk.X, expand=True)
                
                tk.Label(
                    row_frame,
                    text=filename,
                    font=("Microsoft YaHei UI", 10),
                    bg=row_frame['bg'],
                    fg=self.colors['text_muted'],
                    anchor=tk.W
                ).pack(side=tk.LEFT, padx=20, pady=5, fill=tk.X, expand=True)
            
            canvas.pack(side="left", fill="both", expand=True)
            scrollbar.pack(side="right", fill="y")
            
            # 底部按钮
            button_frame = tk.Frame(mapping_window, bg=self.colors['bg_main'])
            button_frame.pack(fill=tk.X, padx=20, pady=(0, 20))
            
            close_btn = UnifiedButton(
                button_frame,
                text="关闭",
                command=mapping_window.destroy,
                style="secondary",
                width=100,
                height=35
            )
            close_btn.pack(side=tk.RIGHT)
            
            # 统计信息
            info_label = tk.Label(
                button_frame,
                text=f"共 {len(mappings)} 个映射",
                font=("Microsoft YaHei UI", 10),
                bg=self.colors['bg_main'],
                fg=self.colors['text_muted']
            )
            info_label.pack(side=tk.LEFT)
            
        except ImportError:
            messagebox.showerror("错误", "无法导入螺旋桨配置模块")
        except Exception as e:
            messagebox.showerror("错误", f"显示映射列表失败：{e}")
    
    def import_pld_file_direct(self, label_type):
        """直接导入PLD文件到指定标签类型目录"""
        try:
            # 根据标签类型确定目标目录
            base_template_dir = self.get_template_directory()
            
            if not base_template_dir:
                # 如果找不到基础模板目录，让用户选择
                base_template_dir = filedialog.askdirectory(title="选择模板根目录")
                if not base_template_dir:
                    return
                base_template_dir = Path(base_template_dir)
            
            # 创建标签类型子目录
            if label_type == "3C":
                target_dir = base_template_dir / "3C标签"
            else:  # 玩具
                target_dir = base_template_dir / "玩具标签"
            
            # 确保目录存在
            target_dir.mkdir(parents=True, exist_ok=True)
            
            # 选择要导入的PLD文件
            file_paths = filedialog.askopenfilenames(
                title=f"选择要导入到{label_type}标签的PLD文件",
                filetypes=[
                    ("PLD文件", "*.pld"),
                    ("所有文件", "*.*")
                ],
                initialdir=os.path.expanduser("~")
            )
            
            if not file_paths:
                return
            
            # 复制文件
            imported_count = 0
            for file_path in file_paths:
                try:
                    import shutil
                    filename = os.path.basename(file_path)
                    target_path = target_dir / filename
                    
                    # 检查文件是否已存在
                    if target_path.exists():
                        result = messagebox.askyesno(
                            "文件已存在",
                            f"文件 {filename} 已存在于{label_type}标签目录中，是否覆盖？"
                        )
                        if not result:
                            continue
                    
                    # 复制文件
                    shutil.copy2(file_path, target_path)
                    imported_count += 1
                    self.log_message(f"✓ 已导入{label_type}标签文件：{filename}")
                    
                except Exception as e:
                    self.log_message(f"✗ 导入文件 {os.path.basename(file_path)} 失败：{e}")
            
            if imported_count > 0:
                messagebox.showinfo("导入完成", f"成功导入 {imported_count} 个PLD文件到{label_type}标签目录")
                
                # 询问是否立即扫描并添加映射
                result = messagebox.askyesno(
                    "添加映射",
                    "是否立即扫描新导入的文件并添加映射？"
                )
                if result:
                    self.auto_scan_and_map()
            else:
                messagebox.showwarning("提示", "没有成功导入任何文件")
                
        except Exception as e:
            self.log_message(f"✗ 导入{label_type}标签文件失败：{e}")
            messagebox.showerror("错误", f"导入{label_type}标签文件失败：{e}")
    
    
    def view_mappings(self):
        """查看所有螺旋桨映射"""
        try:
            from propeller_config import get_all_mappings
            mappings = get_all_mappings()
            
            if not mappings:
                messagebox.showinfo("映射列表", "当前没有螺旋桨映射")
                return
            
            # 创建查看窗口
            view_window = tk.Toplevel(self.root)
            view_window.title("螺旋桨映射列表")
            view_window.geometry("600x400")
            view_window.configure(bg=self.colors['bg_main'])
            
            # 居中显示
            view_window.transient(self.root)
            view_window.grab_set()
            
            # 标题
            title_label = tk.Label(
                view_window,
                text=f"螺旋桨映射列表 (共 {len(mappings)} 个)",
                font=("Microsoft YaHei UI", 16, "bold"),
                bg=self.colors['bg_main'],
                fg=self.colors['text_primary']
            )
            title_label.pack(pady=20)
            
            # 创建文本区域显示映射
            text_frame = tk.Frame(view_window, bg=self.colors['bg_main'])
            text_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 20))
            
            # 文本框和滚动条
            text_widget = scrolledtext.ScrolledText(
                text_frame,
                font=("Microsoft YaHei UI", 10),
                bg=self.colors['bg_card'],
                fg=self.colors['text_primary'],
                relief=tk.FLAT,
                bd=1,
                wrap=tk.WORD
            )
            text_widget.pack(fill=tk.BOTH, expand=True)
            
            # 添加映射内容
            content = "商品编码 -> PLD文件名\n"
            content += "=" * 50 + "\n\n"
            
            for i, (code, filename) in enumerate(sorted(mappings.items()), 1):
                content += f"{i:2d}. {code} -> {filename}\n"
            
            text_widget.insert(tk.END, content)
            text_widget.config(state=tk.DISABLED)  # 设为只读
            
            # 底部按钮
            button_frame = tk.Frame(view_window, bg=self.colors['bg_main'])
            button_frame.pack(fill=tk.X, padx=20, pady=(0, 20))
            
            close_btn = UnifiedButton(
                button_frame,
                text="关闭",
                command=view_window.destroy,
                style="secondary",
                width=100,
                height=35
            )
            close_btn.pack(side=tk.RIGHT)
            
        except ImportError:
            messagebox.showerror("错误", "无法导入螺旋桨配置模块")
        except Exception as e:
            messagebox.showerror("错误", f"查看映射列表失败：{e}")


def main():
    """主函数"""
    try:
        root = tk.Tk()
        app = IntegratedApp(root)
        app.center_window()
        root.mainloop()
    except Exception as e:
        import traceback
        error_msg = f"启动失败：{e}\n\n{traceback.format_exc()}"
        print(error_msg)
        # 显示错误对话框
        try:
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror("标签箱唛工具启动失败", error_msg)
        except:
            pass


if __name__ == "__main__":
    main()

